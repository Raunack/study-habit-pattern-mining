from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def create_report():
    doc = Document()

    # Title
    title = doc.add_heading('Student Study Habit Pattern Mining\nProject Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_paragraph('Date: December 11, 2025').alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph('Submitted by: [Your Name]').alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    doc.add_page_break()

    # Executive Summary
    doc.add_heading('1. Executive Summary', level=1)
    p = doc.add_paragraph(
        'This project analyzes student study habits to predict academic performance using Machine Learning. '
        'It groups students into behavioral "personas" (Clustering) and predicts exam scores (Regression). '
        'The solution is deployed as an interactive Streamlit Web Application.'
    )
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Methodology
    doc.add_heading('2. Methodology & Architecture', level=1)
    doc.add_paragraph(
        'The system processes raw student data through a pipeline of imputation, encoding, and scaling before '
        'feeding it into Random Forest and K-Means models.'
    )
    
    doc.add_heading('2.1 Data Pipeline & Application Flow', level=2)
    
    # Add Diagram Image
    diagram_path = 'images/full_diagrams.png'
    if os.path.exists(diagram_path):
        doc.add_picture(diagram_path, width=Inches(6.5))
        doc.add_paragraph('Figure 1: Data Processing Pipeline and User Flow').alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        doc.add_paragraph('[Error: Diagram images not found]')

    # App Demonstration
    doc.add_page_break()
    doc.add_heading('3. Application Demonstration', level=1)
    doc.add_paragraph('The following screenshots demonstrate the working application.')

    # Screenshots
    screenshots = [
        ('images/app_home.png', 'Figure 2: Application Home Page'),
        ('images/app_results.png', 'Figure 3: Prediction Results'),
        ('images/app_charts.png', 'Figure 4: Comparative Analytics')
    ]

    for path, caption in screenshots:
        if os.path.exists(path):
            doc.add_heading(caption.split(':')[1].strip(), level=2)
            try:
                doc.add_picture(path, width=Inches(6.0))
                doc.add_paragraph(caption).alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.add_paragraph('\n')
            except Exception as e:
                doc.add_paragraph(f"[Error adding image {path}: {e}]")
        else:
             doc.add_paragraph(f"[Image not found: {path}]")

    # Conclusion
    doc.add_heading('4. Conclusion', level=1)
    doc.add_paragraph(
        'The project successfully demonstrates the use of AI in educational analytics. '
        'The web application provides actionable insights to students, helping them optimize their study habits.'
    )

    # Save
    output_path = 'Student_Study_Habit_Project_Report_Fixed.docx'
    doc.save(output_path)
    print(f"Report saved to {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_report()
